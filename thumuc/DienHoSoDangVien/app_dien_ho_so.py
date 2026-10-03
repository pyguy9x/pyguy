# -*- coding: utf-8 -*-
"""
ỨNG DỤNG ĐIỀN TỰ ĐỘNG HỒ SƠ ĐẢNG VIÊN – BẢN THỬ NGHIỆM (5 TRƯỜNG ĐẦU)

Cách hoạt động:
  - Ứng dụng chạy độc lập (cửa sổ riêng), tự mở một cửa sổ Chrome/Edge do ứng dụng điều khiển.
  - Người dùng tự đăng nhập Hệ thống CSDL Đảng viên trong cửa sổ đó (ứng dụng KHÔNG lưu mật khẩu).
  - Mở trang Thêm mới hồ sơ đảng viên, nhập/nạp dữ liệu trong ứng dụng, bấm "Điền vào trang".
  - Ứng dụng KHÔNG bấm Lưu. Người dùng kiểm tra rồi tự bấm "Lưu & tiếp tục".
"""
import os
import queue
import sys
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from dien_tu_dong import TRUONG, LUA_CHON, URL_TAO_MOI, dien_ho_so, tim_trang_tao_moi

THU_MUC = os.path.dirname(os.path.abspath(sys.argv[0]))
HO_SO_TRINH_DUYET = os.path.join(THU_MUC, "du_lieu_trinh_duyet")  # giữ phiên đăng nhập


# ---------------------------------------------------------------------------
# Luồng điều khiển trình duyệt (Playwright chỉ chạy trong 1 luồng riêng)
# ---------------------------------------------------------------------------
class BoDieuKhien(threading.Thread):
    def __init__(self, bao_cao):
        super().__init__(daemon=True)
        self.viec = queue.Queue()
        self.bao_cao = bao_cao
        self.context = None

    def run(self):
        from playwright.sync_api import sync_playwright
        pw = sync_playwright().start()
        while True:
            lenh, tham_so, xong = self.viec.get()
            try:
                if lenh == "mo":
                    self._mo(pw)
                elif lenh == "dien":
                    self._dien(tham_so)
            except Exception as e:
                self.bao_cao(f"Lỗi: {str(e).splitlines()[0]}")
            finally:
                if xong:
                    xong()

    def gui(self, lenh, tham_so=None, xong=None):
        self.viec.put((lenh, tham_so, xong))

    def _mo(self, pw):
        if self.context is not None:
            try:
                p = self.context.pages[0] if self.context.pages else self.context.new_page()
                p.bring_to_front()
                self.bao_cao("Trình duyệt đang mở.")
                return
            except Exception:
                self.context = None
        loi_cuoi = None
        for kenh in ("chrome", "msedge"):
            try:
                self.context = pw.chromium.launch_persistent_context(
                    HO_SO_TRINH_DUYET, channel=kenh, headless=False,
                    no_viewport=True, args=["--start-maximized"])
                break
            except Exception as e:
                loi_cuoi = e
        if self.context is None:
            raise RuntimeError(f"Không mở được Chrome hoặc Edge trên máy: {loi_cuoi}")
        page = self.context.pages[0] if self.context.pages else self.context.new_page()
        page.goto(URL_TAO_MOI)
        self.bao_cao("Đã mở trình duyệt. Hãy đăng nhập và mở trang Thêm mới hồ sơ đảng viên.")

    def _dien(self, du_lieu):
        if self.context is None:
            raise RuntimeError("Chưa mở trình duyệt. Bấm \"1. Mở trình duyệt\" trước.")
        page = tim_trang_tao_moi(self.context)
        if page is None:
            raise RuntimeError("Không thấy trang Thêm mới hồ sơ đảng viên trong cửa sổ trình duyệt của ứng dụng.")
        ten = du_lieu.get("_ten", "")
        self.bao_cao(f"Bắt đầu điền{(' hồ sơ: ' + ten) if ten else ''} ...")
        ok, loi = dien_ho_so(page, du_lieu, self.bao_cao)
        self.bao_cao(f"Xong: {ok} trường đã điền, {len(loi)} lỗi. KIỂM TRA LẠI rồi tự bấm Lưu.")


# ---------------------------------------------------------------------------
# Giao diện
# ---------------------------------------------------------------------------
class UngDung(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Điền tự động hồ sơ đảng viên – bản thử nghiệm")
        self.geometry("760x620")
        self.minsize(680, 560)
        self.thong_bao = queue.Queue()
        self.dong_excel = []
        self.bdk = BoDieuKhien(self.thong_bao.put)
        self.bdk.start()
        self._ve()
        self.after(150, self._nhan_thong_bao)

    def _ve(self):
        pad = {"padx": 10, "pady": 4}
        khung_nut = ttk.Frame(self)
        khung_nut.pack(fill="x", **pad)
        ttk.Button(khung_nut, text="1. Mở trình duyệt", command=self._mo).pack(side="left")
        ttk.Button(khung_nut, text="Nạp dữ liệu từ Excel…", command=self._nap_excel).pack(side="left", padx=8)
        ttk.Button(khung_nut, text="Tạo tệp Excel mẫu…", command=self._tao_mau).pack(side="left")

        khung_ds = ttk.LabelFrame(self, text="Danh sách hồ sơ từ Excel (chọn 1 dòng để đưa lên biểu mẫu)")
        khung_ds.pack(fill="x", **pad)
        self.ds = tk.Listbox(khung_ds, height=5, exportselection=False)
        self.ds.pack(fill="x", padx=6, pady=6)
        self.ds.bind("<<ListboxSelect>>", self._chon_dong)

        khung_f = ttk.LabelFrame(self, text="Dữ liệu sẽ điền (để trống = bỏ qua)")
        khung_f.pack(fill="x", **pad)
        khung_f.columnconfigure(1, weight=1)
        self.o = {}
        for r, (key, nhan, kieu) in enumerate(TRUONG):
            ttk.Label(khung_f, text=nhan).grid(row=r, column=0, sticky="w", padx=6, pady=4)
            if key in LUA_CHON:
                w = ttk.Combobox(khung_f, values=[""] + LUA_CHON[key])
            else:
                w = ttk.Entry(khung_f)
            w.grid(row=r, column=1, sticky="ew", padx=6, pady=4)
            self.o[key] = w
        ttk.Label(khung_f, foreground="#666",
                  text="Tổ chức đảng: gõ tên hoặc mã như trong hộp chọn, ví dụ \"Chi bộ các Cơ quan Đảng\" hoặc \"15.CS023.CB002\".") \
            .grid(row=len(TRUONG), column=0, columnspan=2, sticky="w", padx=6, pady=(0, 6))

        khung_dien = ttk.Frame(self)
        khung_dien.pack(fill="x", **pad)
        self.nut_dien = ttk.Button(khung_dien, text="2. Điền vào trang", command=self._dien)
        self.nut_dien.pack(side="left")
        ttk.Button(khung_dien, text="Xóa ô nhập", command=self._xoa).pack(side="left", padx=8)

        khung_log = ttk.LabelFrame(self, text="Nhật ký")
        khung_log.pack(fill="both", expand=True, **pad)
        self.log = tk.Text(khung_log, height=10, state="disabled", wrap="word")
        self.log.pack(fill="both", expand=True, padx=6, pady=6)
        self._ghi("Bước 1: bấm \"Mở trình duyệt\", đăng nhập, mở trang Thêm mới hồ sơ đảng viên.\n"
                  "Bước 2: nhập hoặc nạp dữ liệu, bấm \"Điền vào trang\". Ứng dụng không tự bấm Lưu.")

    # ---- tiện ích ----
    def _ghi(self, s):
        self.log.configure(state="normal")
        self.log.insert("end", s + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _nhan_thong_bao(self):
        while not self.thong_bao.empty():
            self._ghi(self.thong_bao.get())
        self.after(150, self._nhan_thong_bao)

    def _xoa(self):
        for w in self.o.values():
            w.delete(0, "end")

    # ---- hành động ----
    def _mo(self):
        self._ghi("Đang mở trình duyệt…")
        self.bdk.gui("mo")

    def _dien(self):
        du_lieu = {k: w.get() for k, w in self.o.items()}
        if not any(str(v).strip() for v in du_lieu.values()):
            messagebox.showinfo("Chưa có dữ liệu", "Hãy nhập ít nhất một trường.")
            return
        self.nut_dien.state(["disabled"])
        self.bdk.gui("dien", du_lieu, xong=lambda: self.after(0, lambda: self.nut_dien.state(["!disabled"])))

    def _tao_mau(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        duong_dan = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="mau_ho_so_dang_vien.xlsx",
                                                 filetypes=[("Excel", "*.xlsx")])
        if not duong_dan:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "HoSo"
        ws.append([nhan for _, nhan, _ in TRUONG])
        ws.append([k for k, _, _ in TRUONG])
        for c in ws[1]:
            c.font = Font(bold=True)
        for c in ws[2]:
            c.font = Font(italic=True, color="888888")
        for col, w in zip("ABCDE", (24, 40, 40, 16, 18)):
            ws.column_dimensions[col].width = w
        wb.save(duong_dan)
        self._ghi(f"Đã tạo tệp mẫu: {duong_dan} (dòng 1: tên trường, dòng 2: mã trường – không xóa; dữ liệu từ dòng 3).")

    def _nap_excel(self):
        from openpyxl import load_workbook
        duong_dan = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if not duong_dan:
            return
        ws = load_workbook(duong_dan, data_only=True).active
        hang = list(ws.iter_rows(values_only=True))
        if len(hang) < 2:
            messagebox.showerror("Lỗi", "Tệp không đúng mẫu.")
            return
        ma = [str(x or "").strip() for x in hang[1]]
        khoa = {k for k, _, _ in TRUONG}
        if not khoa & set(ma):
            messagebox.showerror("Lỗi", "Dòng 2 của tệp phải là mã trường (dùng nút \"Tạo tệp Excel mẫu\").")
            return
        self.dong_excel = []
        self.ds.delete(0, "end")
        for h in hang[2:]:
            if not any(v not in (None, "") for v in h):
                continue
            d = {ma[i]: ("" if v is None else str(v)) for i, v in enumerate(h) if i < len(ma) and ma[i] in khoa}
            self.dong_excel.append(d)
            self.ds.insert("end", f"{len(self.dong_excel)}. Số LL: {d.get('profileNumber','')} | Thẻ: {d.get('cardNumber','')} | {d.get('activityOfficialOrgId','')}")
        self._ghi(f"Đã nạp {len(self.dong_excel)} hồ sơ từ {os.path.basename(duong_dan)}.")

    def _chon_dong(self, _e=None):
        sel = self.ds.curselection()
        if not sel:
            return
        d = self.dong_excel[sel[0]]
        self._xoa()
        for k, w in self.o.items():
            w.insert(0, d.get(k, ""))


if __name__ == "__main__":
    UngDung().mainloop()
